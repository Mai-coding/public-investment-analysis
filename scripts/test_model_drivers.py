#!/usr/bin/env python3
"""Disclosure-based model choice and focused scenario coverage regressions."""
from pathlib import Path
import tempfile
import openpyxl

from test_artifact_audit import build_model, run_audit, set_value, source_comment, derived_comment, BLUE


def percentage_model(path):
    build_model(path)
    w = openpyxl.load_workbook(path)
    r = w['Revenue Model']
    r.delete_rows(4, 2)
    r['C2'] = "='Calculation Support'!B3*(1+C3)"
    r['D2'] = derived_comment('Model=GrowthPercent; previous revenue times one plus growth; quantity/price not disclosed')
    r['C3'] = "='Scenarios'!B7"
    support = w['Calculation Support']
    support['B2'] = "='Revenue Model'!C2"
    set_value(support['A3'], 'Prior period revenue')
    set_value(support['B3'], 5000, font=BLUE)
    set_value(support['C3'], source_comment('reported prior-period revenue'))
    s = w['Scenarios']
    s['A7'] = 'Active growth'
    s['B8'], s['B9'], s['B10'] = .05, .1, .15
    for row in range(7,11):
        s.cell(row,2).number_format = '0.0%'
    s['A14'], s['B14'] = 'GROWTH', 'Revenue growth'
    s['H14'] = 'Revenue Model!C3'
    w.save(path)


def add_driver(w, index):
    s = w['Scenarios']; active = 20 + index * 5
    for row, value in ((active, f'=CHOOSE($B$1,B{active+1},B{active+2},B{active+3})'),
                       (active+1,.05),(active+2,.1),(active+3,.15)):
        set_value(s.cell(row,2), value, number_format='0.0%')
    output_row = index + 1
    set_value(w['Income Statement'].cell(output_row,1), f'Driver {index}')
    set_value(w['Income Statement'].cell(output_row,2), f"='Scenarios'!B{active}")
    values = (f'DRIVER-{index}', f'Driver {index}', 'Material and uncertain profit contribution',
              f'Scenarios!B{active}', f'Scenarios!B{active+1}', f'Scenarios!B{active+2}',
              f'Scenarios!B{active+3}', f'Income Statement!B{output_row}')
    for col, value in enumerate(values,1):
        set_value(s.cell(14+index,col),value)


def main():
    with tempfile.TemporaryDirectory() as directory:
        root=Path(directory); base=root/'percent.xlsx'; percentage_model(base)
        result=run_audit(base)
        assert result.returncode == 0, result.stdout
        share=openpyxl.load_workbook(base)
        share['Revenue Model']['A3']='Revenue share %'
        share['Revenue Model']['C2']="='Calculation Support'!B3*C3"
        share['Revenue Model']['D2']=derived_comment('Model=RevenueShare; independently modelled parent revenue times disclosed share assumption')
        path=root/'share.xlsx';share.save(path)
        result=run_audit(path)
        assert result.returncode == 0,result.stdout
        five = openpyxl.load_workbook(base)
        for i in range(1,5): add_driver(five,i)
        path=root/'five.xlsx';five.save(path)
        result=run_audit(path)
        assert result.returncode == 0, result.stdout
        add_driver(five,5);path=root/'six.xlsx';five.save(path)
        result=run_audit(path)
        assert result.returncode and '1–5 key drivers' in result.stdout,result.stdout

        qp=root/'quantity.xlsx';build_model(qp)
        for name,mutate,message in (
            ('undisclosed_qp',lambda w:setattr(w['Revenue Model']['D2'],'value',derived_comment('Model=QuantityPrice; quantity times price')), 'quantity/price Disclosure'),
            ('missing_model',lambda w:setattr(w['Revenue Model']['D2'],'value',derived_comment('quantity times price')), 'must declare Model='),
            ('cosmetic_selector',lambda w:setattr(w['Scenarios']['B7'],'value',100), 'Active input must select'),
            ('unlinked_driver',lambda w:setattr(w['Revenue Model']['C4'],'value',100), 'must consume its Active input'),
            ('unregistered_case',lambda w:set_value(w['Scenarios']['B60'],'=CHOOSE($B$1,80,100,120)'), 'unregistered scenario-selecting input'),
            ('large_font',lambda w:setattr(w['Revenue Model']['C2'],'font',openpyxl.styles.Font(name='Arial',size=14)), 'model font size must be 8–10'),
            ('gridlines',lambda w:setattr(w['Revenue Model'].sheet_view,'showGridLines',True), 'hidden gridlines'),
        ):
            w=openpyxl.load_workbook(qp);mutate(w);path=root/f'{name}.xlsx';w.save(path)
            result=run_audit(path)
            assert result.returncode and message in result.stdout,(name,result.stdout)
    print('MODEL DRIVER AND FORMAT TESTS: PASS (percentage/disclosed QP, 1/5/6 drivers, linkage and layout)')
    return 0


if __name__=='__main__':
    raise SystemExit(main())
