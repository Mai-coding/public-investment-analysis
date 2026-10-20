"""Exercise presentation requirements on real DOCX files and effective inheritance."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from report_content_audit import Audit, audit_cover_tables
from report_layout_rules import (audit_image_paragraphs, audit_report_tables,
                                 audit_toc_page_colors, page_number_runs)
from test_report_content_audit import PNG, build


def xml_color(run, value='061F32', theme=None):
    props = run.find(qn('w:rPr'))
    if props is None:
        props = OxmlElement('w:rPr'); run.insert(0, props)
    for previous in list(props.findall(qn('w:color'))):
        props.remove(previous)
    color = OxmlElement('w:color'); color.set(qn('w:val'), value)
    if theme:
        color.set(qn('w:themeColor'), theme)
    props.append(color)


class ReportStyleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        image = cls.root / 'pixel.png'; image.write_bytes(PNG)
        cls.baseline = cls.root / 'report.docx'; build(cls.baseline, image)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def mutate(self, operation):
        doc = Document(self.baseline)
        operation(doc)
        path = self.root / 'changed.docx'; doc.save(path)
        return path

    def check(self, function, path, failure=None):
        audit = Audit(); function(path, audit)
        if failure:
            self.assertTrue(any(failure in item for item in audit.failures), audit.failures)
        else:
            self.assertEqual(audit.failures, [])

    def image_paragraph(self, doc):
        return next(p for p in doc.paragraphs if p._p.find('.//' + qn('w:drawing')) is not None)

    def toc(self, doc):
        return next(p for p in doc.paragraphs if p.style.name == 'ReportTOC')

    def test_baseline_all_requirements(self):
        for function in (audit_cover_tables, audit_report_tables, audit_toc_page_colors, audit_image_paragraphs):
            self.check(function, self.baseline)

    def test_cover_old_columns_rejected(self):
        path = self.mutate(lambda d: setattr(d.tables[0].cell(0, 0), 'text', 'Current band'))
        self.check(audit_cover_tables, path, 'Cover columns must be')

    def test_cover_price_needs_current_and_base(self):
        path = self.mutate(lambda d: setattr(d.tables[0].cell(1, 0), 'text', 'Current price: 57 USD'))
        self.check(audit_cover_tables, path, 'current price and Base-case price')

    def test_return_excludes_other_scenarios_and_entry(self):
        for value in ('Base 12%; Bear -24%', 'Base 12%; Attractive entry 45 USD', '12% CAGR'):
            with self.subTest(value=value):
                path = self.mutate(lambda d: setattr(d.tables[0].cell(1, 1), 'text', value))
                self.check(audit_cover_tables, path, 'only the Base-case return')

    def test_body_cell_alignment_is_checked_beyond_cover(self):
        path = self.mutate(lambda d: setattr(d.tables[1].cell(1, 0).paragraphs[0], 'alignment', WD_ALIGN_PARAGRAPH.LEFT))
        self.check(audit_report_tables, path, 'horizontally and vertically centered')

    def test_appendix_vertical_alignment_is_checked(self):
        def change(d):
            table = d.add_table(rows=2, cols=1)
            table.cell(0, 0).text = 'Appendix data'
            table.cell(1, 0).text = 'Item'
            table.cell(1, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        self.check(audit_report_tables, self.mutate(change), 'horizontally and vertically centered')

    def test_header_fill_uses_exact_navy(self):
        def change(d):
            shade = d.tables[1].cell(0, 0)._tc.find('w:tcPr/w:shd', d.tables[1]._tbl.nsmap)
            shade.set(qn('w:fill'), '2F80ED')
        self.check(audit_report_tables, self.mutate(change), 'NAVY/white')

    def test_header_theme_fill_is_not_allowed(self):
        def change(d):
            shade = d.tables[0].cell(0, 0)._tc.find('w:tcPr/w:shd', d.tables[0]._tbl.nsmap)
            shade.set(qn('w:themeFill'), 'accent1')
        self.check(audit_report_tables, self.mutate(change), 'NAVY/white')

    def test_header_white_text_is_checked(self):
        path = self.mutate(lambda d: xml_color(d.tables[1].cell(0, 0).paragraphs[0].runs[0]._r, '061F32'))
        self.check(audit_report_tables, path, 'NAVY/white')

    def test_pure_image_container_has_no_required_data_header(self):
        def change(d):
            table = d.add_table(rows=1, cols=1)
            image = self.root / 'pixel.png'
            table.cell(0, 0).paragraphs[0].add_run().add_picture(str(image))
        self.check(audit_report_tables, self.mutate(change))

    def test_page_color_does_not_confuse_chapter_label_with_page(self):
        path = self.mutate(lambda d: xml_color(self.toc(d)._p.find('.//' + qn('w:hyperlink')).find(qn('w:r')), '2F80ED'))
        self.check(audit_toc_page_colors, path)

    def test_simple_pageref_page_color(self):
        path = self.mutate(lambda d: xml_color(self.toc(d)._p.find('.//' + qn('w:fldSimple')).find(qn('w:r')), '2F80ED'))
        self.check(audit_toc_page_colors, path, 'TOC page-number colors')

    def test_pageref_theme_override(self):
        path = self.mutate(lambda d: xml_color(self.toc(d)._p.find('.//' + qn('w:fldSimple')).find(qn('w:r')), theme='hyperlink'))
        self.check(audit_toc_page_colors, path, 'TOC page-number colors')

    def test_complex_pageref_result_is_detected(self):
        def change(d):
            link = self.toc(d)._p.find('.//' + qn('w:hyperlink'))
            fld = link.find(qn('w:fldSimple')); instruction = fld.get(qn('w:instr')); link.remove(fld)
            for kind in ('begin', 'instruction', 'separate', 'result', 'end'):
                run = OxmlElement('w:r')
                if kind == 'instruction':
                    node = OxmlElement('w:instrText'); node.text = instruction
                elif kind == 'result':
                    node = OxmlElement('w:t'); node.text = '23'; xml_color(run, 'FF0000')
                else:
                    node = OxmlElement('w:fldChar'); node.set(qn('w:fldCharType'), kind)
                run.append(node); link.append(run)
        self.check(audit_toc_page_colors, self.mutate(change), 'TOC page-number colors')

    def test_native_tab_delimited_page(self):
        def change(d):
            p = self.toc(d)._p
            link = p.find('.//' + qn('w:hyperlink'))
            link.remove(link.find(qn('w:fldSimple')))
            run = OxmlElement('w:r'); run.append(OxmlElement('w:tab')); link.append(run)
            run = OxmlElement('w:r'); node = OxmlElement('w:t'); node.text = '17'; run.append(node)
            xml_color(run, 'FF0000'); link.append(run)
            self.assertEqual(len(page_number_runs(p)), 1)
        self.check(audit_toc_page_colors, self.mutate(change), 'TOC page-number colors')

    def test_image_inherited_marker_flags_and_explicit_false(self):
        def inherited(d):
            p = self.image_paragraph(d)
            style = d.styles.add_style('FigureHost', 1)
            style.paragraph_format.keep_with_next = True
            p.style = style
        path = self.mutate(inherited)
        self.check(audit_image_paragraphs, path, 'keepNext')
        doc = Document(path); self.image_paragraph(doc).paragraph_format.keep_with_next = False; doc.save(path)
        self.check(audit_image_paragraphs, path)

    def test_all_known_image_paragraph_flags(self):
        for setting in ('keep_together', 'page_break_before'):
            with self.subTest(setting=setting):
                path = self.mutate(lambda d: setattr(self.image_paragraph(d).paragraph_format, setting, True))
                self.check(audit_image_paragraphs, path, 'square-marker causes')

    def test_image_list_cancellation_and_literal_square(self):
        def change(d):
            p = self.image_paragraph(d)
            props = p._p.get_or_add_pPr(); num = OxmlElement('w:numPr'); item = OxmlElement('w:numId')
            item.set(qn('w:val'), '1'); num.append(item); props.append(num)
        path = self.mutate(change)
        self.check(audit_image_paragraphs, path, 'list numbering')
        doc = Document(path); p = self.image_paragraph(doc)
        p._p.find('w:pPr/w:numPr/w:numId', p._p.nsmap).set(qn('w:val'), '0'); doc.save(path)
        self.check(audit_image_paragraphs, path)
        path = self.mutate(lambda d: self.image_paragraph(d).add_run('■'))
        self.check(audit_image_paragraphs, path, 'literal square')

    def test_unrelated_heading_keep_with_next_is_preserved(self):
        path = self.mutate(lambda d: setattr(next(p for p in d.paragraphs if p.style.name == 'Heading 1').paragraph_format, 'keep_with_next', True))
        self.check(audit_image_paragraphs, path)


if __name__ == '__main__':
    unittest.main()
