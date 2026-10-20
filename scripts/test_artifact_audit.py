#!/usr/bin/env python3
"""Regression tests for the public-investment model artifact audit."""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from docx import Document
from openpyxl.styles import Font

from terminal_multiple import TerminalCase


HERE = Path(__file__).resolve().parent
AUDIT = HERE / "artifact_audit.py"
ARIAL = Font(name="Arial", size=8, color="FF000000")
BLUE = Font(name="Arial", size=8, color="FF0000FF")


def set_value(cell, value, font=ARIAL, number_format=None) -> None:
    cell.value = value
    cell.font = font
    if number_format:
        cell.number_format = number_format


def source_comment(basis: str) -> str:
    return (
        "Class=Fact; Source=Test Company 2025 Form 10-K; Date=2026-02-20; "
        f"Locator=Item 8, p.10; Basis={basis}"
    )


def derived_comment(basis: str) -> str:
    return f"Class=Derived; Source=N/A; Date=N/A; Locator=N/A; Basis={basis}"


def build_model(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in (
        "Revenue Model",
        "Calculation Support",
        "Income Statement",
        "Cash Flow Statement",
        "Balance Sheet",
        "Scenarios",
        "Valuation Inputs",
        "Comps",
        "Terminal Cases",
        "Valuation Summary",
    ):
        wb.create_sheet(name)

    ws = wb["Revenue Model"]
    for col, value in enumerate(("Metric", "Unit", "2027E", "Comments"), 1):
        set_value(ws.cell(1, col), value)
    rows = (
        ("Energy revenue", "$M", "=C4*C5", derived_comment("Model=QuantityPrice; quantity times price; Disclosure=Test Company 2025 Form 10-K Item 8 p.10, reported volume and realized price"), ARIAL, "$0.0"),
        ("YoY growth", "%", "=10%", derived_comment("current year revenue divided by prior year revenue minus one"), ARIAL, "0.0%"),
        ("Quantity", "MWh", "='Scenarios'!B7", source_comment("forecast volume calibrated to disclosed capacity"), ARIAL, "0.0"),
        ("Price / yield", "$/MWh", 50, source_comment("forecast price calibrated to disclosed realized price"), BLUE, "$0.0"),
    )
    for row, (metric, unit, value, comment, font, fmt) in enumerate(rows, 2):
        set_value(ws.cell(row, 1), metric)
        set_value(ws.cell(row, 2), unit)
        set_value(ws.cell(row, 3), value, font=font, number_format=fmt)
        set_value(ws.cell(row, 4), comment)

    support = wb["Calculation Support"]
    for col, value in enumerate(("Metric", "2027E", "Comments"), 1):
        set_value(support.cell(1, col), value)
    set_value(support["A2"], "Volume sanity check")
    set_value(support["B2"], "='Revenue Model'!C4")
    set_value(support["C2"], derived_comment("links displayed quantity to the physical sanity check"))

    income = wb["Income Statement"]
    set_value(income["A1"], "Revenue")
    set_value(income["B1"], "='Revenue Model'!C2")

    cash = wb["Cash Flow Statement"]
    for row, label in enumerate(("Maintenance CapEx", "Growth CapEx", "Acquisition / Integration CapEx"), 1):
        set_value(cash.cell(row, 1), label)
        set_value(cash.cell(row, 2), 1)

    balance = wb["Balance Sheet"]
    set_value(balance["A1"], "Balance Check")
    set_value(balance["B1"], "=0")

    scenario = wb['Scenarios']
    set_value(scenario['A1'], 'Case Selector')
    set_value(scenario['B1'], 2, font=BLUE)
    for row, (label, value) in enumerate((('Active volume', '=CHOOSE($B$1,B8,B9,B10)'), ('Bear', 80), ('Base', 100), ('Bull', 120)), 7):
        set_value(scenario.cell(row, 1), label)
        set_value(scenario.cell(row, 2), value)
    set_value(scenario['A12'], 'Scenario Drivers')
    for col, value in enumerate(('Driver ID', 'Driver', 'Reason', 'Active Ref', 'Bear Ref', 'Base Ref', 'Bull Ref', 'Output Ref'), 1):
        set_value(scenario.cell(13, col), value)
    for col, value in enumerate(('VOLUME', 'Volume', 'Material revenue uncertainty', 'Scenarios!B7', 'Scenarios!B8', 'Scenarios!B9', 'Scenarios!B10', 'Revenue Model!C4'), 1):
        set_value(scenario.cell(14, col), value)
    set_value(wb['Valuation Inputs']['A1'], 'Valuation Inputs')

    comps = wb["Comps"]
    set_value(comps["A1"], "Company")
    set_value(comps["B1"], "Role")
    for row, company in enumerate(("Peer A", "Peer B", "Peer C"), 2):
        set_value(comps.cell(row, 1), company)
        set_value(comps.cell(row, 2), "Central primary")

    valuation = wb["Valuation Summary"]
    set_value(valuation["A1"], "Valuation Method Manifest")
    headers = ("Method ID", "Method Label", "Role", "Low", "Mid", "High", "Weight", "Source Range / Formula", "Notes", "Method Family")
    for col, value in enumerate(headers, 1):
        set_value(valuation.cell(2, col), value)
    methods = (
        ("VAL-PE", "2027E Forward P/E", "Central", 90, 100, 110, 1.0, "Comps!A2:B4", "next-year EPS", "Forward Comps"),
        ("VAL-TERM-PE", "Terminal P/E reference", "Diagnostic", 85, 100, 115, 0.0, "Terminal Cases!A3:O3", "operating continuation", "Terminal"),
    )
    for row, values in enumerate(methods, 3):
        for col, value in enumerate(values, 1):
            set_value(valuation.cell(row, col), value, number_format="0.0%" if col == 7 else None)
    set_value(valuation["A7"], "Price Band Map")
    band_headers = (
        "Band", "Low", "High", "Operating State", "Valuation Regime", "Required Conditions",
        "Key KPIs", "Driver (E / Multiple)", "Distance", "Robustness", "Actionability",
        "Evidence Status", "Evidence Ref",
    )
    for col, value in enumerate(band_headers, 1):
        set_value(valuation.cell(8, col), value)
    bands = (
        ("Upside / Re-rating", 125, 150, "EPS +25%", "P/E 30-33x", "new product ramps", "NRR>120%", "Both", "High", "Low", "Trim"),
        ("Base+", 108, 125, "EPS +18%", "P/E 27-30x", "margin beat", "GM>70%", "E", "Medium", "Medium", "Hold"),
        ("Base", 90, 108, "EPS +12%", "P/E 24-27x", "guide met", "NRR~112%", "E", "Low", "High", "Initiate"),
        ("Disappointment", 72, 90, "EPS +6%", "P/E 20-23x", "one guide cut", "NRR<108%", "Multiple", "Medium", "Medium", "Accumulate"),
        ("Thesis Impairment", 50, 72, "EPS flat", "P/E 16-19x", "share loss", "NRR<100%", "Both", "High", "Low (stacked)", "Conditional"),
    )
    for row, values in enumerate(bands, 9):
        for col, value in enumerate(values, 1):
            set_value(valuation.cell(row, col), value)
        set_value(valuation.cell(row, 12), "Supported" if values[0] == 'Base' else 'Partial')
        set_value(valuation.cell(row, 13), f"Condition Evidence row {row + 21}")
    set_value(valuation["A15"], "3Y Return Bridge")
    for row, (label, value) in enumerate((
        ("EPS per-share CAGR", 0.12), ("Shareholder yield (dividend only)", 0.01),
        ("Annualized multiple change", -0.02), ("Expected 3Y CAGR", 0.11),
    ), 16):
        set_value(valuation.cell(row, 1), label)
        set_value(valuation.cell(row, 2), value, number_format="0.0%")

    set_value(valuation['A22'], 'Decision Position')
    for col, value in enumerate(('Position', 'Band', 'Confidence', 'Rationale'), 1):
        set_value(valuation.cell(23, col), value)
    for row, values in enumerate((
        ('Market-priced Band', 'Base', 'High', 'Current price is inside the Base price interval'),
        ('Evidence-supported Band', 'Base', 'Medium', 'Retention supports Base; product ramp not yet verified'),
    ), 24):
        for col, value in enumerate(values, 1):
            set_value(valuation.cell(row, col), value)
    set_value(valuation['A28'], 'Condition Evidence')
    columns = ('Band', 'Condition', 'Threshold', 'Current Evidence', 'Evidence Date',
               'Source / Locator', 'Status', 'Decisive', 'Next Signal')
    for col, value in enumerate(columns, 1):
        set_value(valuation.cell(29, col), value)
    for row, band in enumerate(bands, 30):
        values = (band[0], band[5], 'Verify KPI over next two quarters', 'Dated retention disclosures',
                  '2026-10-17', 'Test Company 2026 Q2 filing, p.10',
                  'Supported' if band[0] == 'Base' else 'Partial', 'Yes', 'Next quarterly retention report')
        for col, value in enumerate(values, 1):
            set_value(valuation.cell(row, col), value)
    terminal = wb['Terminal Cases']
    set_value(terminal['A1'], 'Terminal Cases')
    columns = ('Case', 'Metric', 'Base Year', 'Current Price', 'Horizon', 'Required Return', 'Exit Multiple',
               'Terminal Total', 'Terminal Shares', 'Distributions FV', 'Supported Price', 'Implied Terminal Total', 'Basis',
               'Terminal Cash', 'Cash Treatment')
    for col, value in enumerate(columns, 1):
        set_value(terminal.cell(2, col), value)
    case = TerminalCase('P/E', 60, 10, 20, 2, .1, 1)
    values = ('Base', case.metric, 2027, 100, case.horizon, case.required_return, case.exit_multiple,
              case.terminal_total, case.terminal_shares, case.distributions_fv, case.supported_price,
              case.implied_total(100), source_comment('Operating continuation with cash included'), 0, 'Included')
    for col, value in enumerate(values, 1):
        set_value(terminal.cell(3, col), value)
    set_value(valuation['A37'], 'Band Valuation Bridge')
    columns = ('Case', 'Comps Low', 'Comps Mid', 'Comps High', 'Terminal Low', 'Terminal Mid',
               'Terminal High', 'Central Low', 'Central Mid', 'Central High')
    for col, value in enumerate(columns, 1):
        set_value(valuation.cell(38, col), value)
    for row, band in enumerate(bands, 39):
        low, high = band[1:3]
        mid = (low + high) / 2
        values = (band[0], low, mid, high, low*.7, mid*.7, high*.7, low, mid, high)
        for col, value in enumerate(values, 1):
            set_value(valuation.cell(row, col), value)
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'C2'
        sheet.column_dimensions['A'].width = 35
        sheet.column_dimensions['B'].width = 10
        for row in sheet:
            sheet.row_dimensions[row[0].row].height = 13.5
        for col in range(3, sheet.max_column + 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 10
    wb['Revenue Model'].column_dimensions['D'].width = 80
    wb['Calculation Support'].column_dimensions['C'].width = 80
    wb.save(path)


def run_audit(model: Path, report: Path | None = None) -> subprocess.CompletedProcess[str]:
    command = [sys.executable, '-X', 'utf8', str(AUDIT), "--excel", str(model)]
    if report:
        command.extend(("--report", str(report)))
    return subprocess.run(command, capture_output=True, text=True, encoding="utf-8")


def main() -> int:
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        model = root / "positive.xlsx"
        report = root / "report.docx"
        build_model(model)
        doc = Document()
        doc.add_paragraph("2027E Forward P/E100%; Terminal P/E reference0%, diagnostic only")
        doc.save(report)
        positive = run_audit(model, report)
        if positive.returncode != 0:
            print(positive.stdout)
            raise AssertionError("positive fixture must pass")

        operating = openpyxl.load_workbook(model)
        for sheet in ('Comps', 'Valuation Summary', 'Terminal Cases'):
            operating.remove(operating[sheet])
        operating_path = root / 'task2_operating_only.xlsx'
        operating.save(operating_path)
        result = run_audit(operating_path)
        assert result.returncode == 0, result.stdout

        negative = root / "negative.xlsx"
        wb = openpyxl.load_workbook(model)
        wb.create_sheet("Assumption Guide", 0)
        wb["Assumption Guide"]["A1"] = "deprecated"
        wb["Assumption Guide"]["A1"].font = ARIAL
        wb["Calculation Support"].sheet_state = "hidden"
        wb["Revenue Model"]["D4"] = "company commentary"
        wb["Revenue Model"]["D4"].font = ARIAL
        wb["Valuation Summary"]["B4"] = "2028E Forward P/S"
        wb["Valuation Summary"]["B4"].font = ARIAL
        bands_ws = wb["Valuation Summary"]
        bands_ws["I9"] = "Far"
        for row, (low, high) in enumerate(((125, 150), (100, 125), (75, 100), (50, 75), (20, 50)), 9):
            bands_ws.cell(row, 2).value = low
            bands_ws.cell(row, 3).value = high
        wb.save(negative)
        failed = run_audit(negative)
        combined = failed.stdout + failed.stderr
        expected = (
            "remove deprecated governance sheet",
            "Calculation Support must be visible",
            "missing field(s)",
            "forward multiples use mixed forecast years",
            "Distance to Band must be Low/Medium/High",
            "mechanically equal-width",
        )
        if failed.returncode == 0 or any(message not in combined for message in expected):
            print(combined)
            raise AssertionError("negative fixture did not trigger every expected rule")

        mutations = (
            ('missing_evidence', lambda w: setattr(w['Valuation Summary']['A28'], 'value', 'Removed'), "missing 'Condition Evidence'"),
            ('incorrect_value', lambda w: setattr(w['Terminal Cases']['K3'], 'value', 999), 'Supported Price disagrees'),
            ('incorrect_inverse', lambda w: setattr(w['Terminal Cases']['L3'], 'value', 999), 'Implied Terminal Total disagrees'),
            ('dilution_omitted', lambda w: setattr(w['Terminal Cases']['I3'], 'value', 12), 'Supported Price disagrees'),
            ('decisive_refutation', lambda w: setattr(w['Valuation Summary']['G32'], 'value', 'Contradicted'), 'contradicted decisive condition'),
            ('missing_position', lambda w: setattr(w['Valuation Summary']['A25'], 'value', 'Removed'), 'missing or incomplete evidence-supported band'),
            ('incorrect_central', lambda w: setattr(w['Valuation Summary']['I41'], 'value', 999), 'must equal100% forward comps'),
            ('old_blend', lambda w: setattr(w['Valuation Summary']['I41'], 'value', .8*w['Valuation Summary']['C41'].value+.2*w['Valuation Summary']['F41'].value), 'must equal100% forward comps'),
            ('terminal_formula_leak', lambda w: setattr(w['Valuation Summary']['I41'], 'value', '=C41+0*F41'), 'depends on a zero-weight terminal reference'),
            ('other_case_terminal_leak', lambda w: setattr(w['Valuation Summary']['I41'], 'value', '=E40'), 'depends on a zero-weight terminal reference'),
            ('indirect_terminal_leak', lambda w: (setattr(w['Valuation Summary']['I41'], 'value', '=N41'), setattr(w['Valuation Summary']['N41'], 'value', "='Terminal Cases'!K3")), 'depends on a zero-weight terminal reference'),
            ('incorrect_weights', lambda w: (setattr(w['Valuation Summary']['G3'], 'value', .8), setattr(w['Valuation Summary']['G4'], 'value', .2)), 'Forward Comps Central weight must be 100%'),
            ('terminal_central_role', lambda w: setattr(w['Valuation Summary']['C4'], 'value', 'Central'), 'Terminal must be Diagnostic'),
            ('cash_double_count', lambda w: setattr(w['Terminal Cases']['N3'], 'value', 100), 'Included cash must not be added again'),
        )
        for name, mutate, expected_message in mutations:
            changed = openpyxl.load_workbook(model)
            mutate(changed)
            path = root / f'{name}.xlsx'
            changed.save(path)
            result = run_audit(path)
            assert result.returncode and expected_message in result.stdout, (name, result.stdout)

        # References do not affect the central value and may be unavailable.
        for name in ('terminal_changed', 'reference_missing', 'separate_cash'):
            changed = openpyxl.load_workbook(model)
            if name == 'reference_missing':
                changed.remove(changed['Terminal Cases'])
                for col in range(5, 8):
                    for row in range(39, 44):
                        changed['Valuation Summary'].cell(row, col).value = 'n.a.'
            else:
                cash = 100 if name == 'separate_cash' else 0
                case = TerminalCase('P/E', 90, 12, 24, 2, .1, 1, cash, 'Separate' if cash else 'Included')
                for addr, value in {'G3':24, 'H3':90, 'I3':12, 'K3':case.supported_price,
                                    'L3':case.implied_total(100), 'N3':cash, 'O3':case.cash_treatment}.items():
                    changed['Terminal Cases'][addr] = value
                for col in range(5, 8):
                    for row in range(39, 44):
                        changed['Valuation Summary'].cell(row, col).value *= 2
            path = root / (name + '.xlsx');changed.save(path)
            result = run_audit(path)
            assert result.returncode == 0, (name, result.stdout)

        # Price and evidence may legitimately disagree: never force matching bands.
        changed = openpyxl.load_workbook(model)
        changed['Valuation Summary']['B24'] = 'Upside / Re-rating'
        path = root / 'market_evidence_disagreement.xlsx'
        changed.save(path)
        result = run_audit(path)
        assert result.returncode == 0, result.stdout
        for metric, total, multiple in (('P/S', 300, 4), ('P/B', 600, 2)):
            changed = openpyxl.load_workbook(model)
            changed['Valuation Summary']['B3'] = f'2027E Forward {metric}'
            changed['Valuation Summary']['B4'] = f'Terminal {metric} reference'
            changed['Terminal Cases']['B3'] = metric
            changed['Terminal Cases']['G3'] = multiple
            changed['Terminal Cases']['H3'] = total
            changed['Terminal Cases']['L3'] = total
            path = root / f'positive_{metric.replace("/", "")}.xlsx'
            changed.save(path)
            result = run_audit(path)
            assert result.returncode == 0, result.stdout

    print("artifact_audit regression tests: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
