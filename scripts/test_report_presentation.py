"""Exercise display behavior and DOCX boundary/asset checks with real documents."""
from pathlib import Path
import tempfile
import zipfile

from docx import Document

from report_presentation import format_number, format_amount, format_percent, style_issues
from report_content_audit import Audit, audit_reader_presentation
from test_report_content_audit import PNG, add_image


def check_formatting():
    for value, expected in [
        (None, '—'), (0, '0'), (6.74, '6.7'), (6, '6'), (9.96, '10'),
        (10.04, '10'), (10.5, '11'), (198.875, '199'), (-23.64, '-24'),
        (-6.74, '-6.7'), (0.004, '<0.1'), (-0.004, '>-0.1'), (11500, '11,500'),
    ]:
        assert format_number(value) == expected, (value, format_number(value))
    for value, unit, expected in [
        (11500000000, None, '115 亿美元'), (1500000000000, None, '15,000 亿美元'),
        (700000000, None, '7 亿美元'), (205000000, None, '2.1 亿美元'),
        (99996000, None, '1 亿美元'), (9999.6, None, '1 万美元'),
        (120000, '亿', '<0.1 亿美元'), (0, '亿', '0 亿美元'),
        (-205000000, None, '-2.1 亿美元'), (None, '亿', '—'),
    ]:
        assert format_amount(value, unit=unit) == expected, (value, unit, format_amount(value, unit=unit))
    assert format_amount(50200, '台') == '5 万台'
    assert format_amount(100000000, '港元') == '1 亿港元'
    assert format_percent(24.36) == '24%'
    assert format_percent(.0674, fraction=True) == '6.7%'
    assert format_percent(.004) == '<0.1%'
    for invalid in ['NaN', 'Infinity', '-Infinity']:
        try:
            format_number(invalid)
        except ValueError:
            pass
        else:
            raise AssertionError('Non-finite value accepted')
    for text in ['增长24.36%', '收入 11.5 亿美元', '收入 115.0 亿美元', '市值 US$11.5bn', '收入 5 百万美元', '统一口径后增长']:
        assert style_issues(text), text
    for text in ['增长24%', '增速6.7%', '收入 115 亿美元', '12.5 Price Band Decision Map', '9.3.1 CEO', '2026.10.16', '2026.10', 'v2.1', '300124.SZ', '1.6T 产品', '3.5万台', '—']:
        assert not style_issues(text), (text, style_issues(text))


def build(path, image, body='增长 6.7%，收入 115 亿美元。', logo=True):
    doc = Document()
    if logo:
        add_image(doc, image, 'COMPANY_LOGO: Test Company')
    doc.add_heading('1. 投资概要', level=1)
    p = doc.add_paragraph()
    # Split a bad decimal across runs: the audit must reconstruct the number.
    for part in [body[:4], body[4:]]:
        p.add_run(part)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '6.7'
    table.cell(0, 1).text = '10'
    doc.add_heading('附录 A 可比公司', level=1)
    doc.add_paragraph('市盈率 24 倍。')
    doc.add_heading('Appendix｜Methods / 定义与方法', level=1)
    doc.add_paragraph('口径：百万美元，未舍入金额 11.52347。')
    doc.add_heading('Appendix｜Sources / 来源索引', level=1)
    doc.add_paragraph('[1] 原始资料标题含 24.36% 和 million，保留原文。')
    doc.save(path)


def audit(path):
    result = Audit()
    audit_reader_presentation(path, result)
    return result


def check_docx():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        image = td / 'logo.png'
        image.write_bytes(PNG)
        path = td / 'sample.docx'
        build(path, image)
        assert not audit(path).failures, audit(path).failures
        for text in ['增长24.36%', '收入 11.5 亿美元', '收入 3 百万美元', '按口径计算']:
            build(path, image, body=text)
            assert any('Reader presentation violations' in f for f in audit(path).failures), text
        build(path, image, logo=False)
        assert any('COMPANY_LOGO' in f for f in audit(path).failures)
        build(path, image)
        # A tagged drawing with a missing embedded asset must fail.
        broken = td / 'broken.docx'
        with zipfile.ZipFile(path) as src, zipfile.ZipFile(broken, 'w') as dest:
            for name in src.namelist():
                if not name.startswith('word/media/'):
                    dest.writestr(name, src.read(name))
        assert any('COMPANY_LOGO' in f for f in audit(broken).failures)
        build(path, image)
        doc = Document(path)
        for p in doc.paragraphs:
            if p.text.startswith('Appendix｜Methods'):
                p.text = '方法记录'
        doc.save(path)
        assert any('exactly one Methods' in f for f in audit(path).failures)


if __name__ == '__main__':
    check_formatting()
    check_docx()
    print('REPORT PRESENTATION TEST: PASS')
