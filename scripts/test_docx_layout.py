#!/usr/bin/env python3
"""Behavioral regressions for real DOCX navigation, inheritance and quadrant evidence."""
from __future__ import annotations

import copy
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from docx_layout_audit import audit_layout
from docx_structure import NS, Structure, chapter_number, text_of
from report_content_audit import (Audit, audit_competition_quadrant, audit_cover_tables,
                                  audit_reader_presentation, chapter_ranges, parse_docx)
from depth_audit import parse_docx as depth_parse, split_sections
from test_report_content_audit import PNG, build


class LayoutTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        image = cls.root / 'pixel.png'; image.write_bytes(PNG)
        cls.baseline = cls.root / 'baseline.docx'; build(cls.baseline, image)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def mutate(self, fn):
        doc = Document(self.baseline); fn(doc)
        path = self.root / 'mutated.docx'; doc.save(path)
        return path

    def check(self, path, message=None, cover=False):
        audit = Audit()
        (audit_cover_tables if cover else audit_layout)(path, audit)
        if message:
            self.assertTrue(any(message in f for f in audit.failures), audit.failures)
        else:
            self.assertEqual(audit.failures, [])

    def test_baseline_native_outline_and_contents(self):
        self.check(self.baseline)
        self.check(self.baseline, cover=True)

    def test_chapter_punctuation_and_chinese_variants(self):
        for title in ('§2 核心技术入门', '§2. 核心技术入门', '2. 核心技术入门', '2 核心技术入门', '第二章 核心技术入门', '第2章 核心技术入门'):
            with self.subTest(title=title):
                path = self.mutate(lambda d: setattr(next(p for p in d.paragraphs if p.text == '2. 核心技术入门'), 'text', title))
                self.assertEqual(chapter_number(title), 2)
                paras = parse_docx(path)
                self.assertEqual(paras[chapter_ranges(paras)[2][0]].text, title)
                chapters, _, _ = split_sections(depth_parse(path))
                self.assertEqual(chapters[2]['title'], title)

    def test_toc_does_not_steal_chapters_or_appendices(self):
        paras = parse_docx(self.baseline)
        self.assertFalse(any(p.text.startswith('01 ') for p in paras))
        self.assertEqual(sum('Appendix｜Methods' in p.text for p in paras), 1)
        a = Audit(); audit_reader_presentation(self.baseline, a)
        self.assertEqual(a.failures, [])
        chapters, _, _ = split_sections(depth_parse(self.baseline))
        self.assertEqual(chapters[2]['title'], '2. 核心技术入门')

    def test_plain_numbered_prose_cannot_start_chapter(self):
        def change(d):
            p = next(p for p in d.paragraphs if p.text == '1. 投资概要')
            p.insert_paragraph_before('2. this is ordinary numbered prose')
        path = self.mutate(change)
        paras = parse_docx(path)
        self.assertEqual(paras[chapter_ranges(paras)[2][0]].text, '2. 核心技术入门')
        chapters, _, _ = split_sections(depth_parse(path))
        self.assertEqual(chapters[2]['title'], '2. 核心技术入门')

    def test_three_row_cover_fails(self):
        self.check(self.mutate(lambda d: d.tables[0].add_row()), 'exactly two rows', cover=True)

    def test_one_row_cover_fails(self):
        self.check(self.mutate(lambda d: d.tables[0]._tbl.remove(d.tables[0].rows[1]._tr)), 'exactly two rows', cover=True)

    def test_missing_native_chapter_does_not_scan_all_body_tables_as_cover(self):
        def change(d):
            next(p for p in d.paragraphs if p.text == '1. 投资概要').style = d.styles['Normal']
        self.check(self.mutate(change), 'Cannot locate native chapter 1', cover=True)

    def test_direct_wrong_color(self):
        self.check(self.mutate(lambda d: setattr(next(p for p in d.paragraphs if p.text == '2.1 技术机制').runs[0].font.color, 'rgb', RGBColor(0, 0, 0))), 'palette mismatch')

    def test_inherited_heading_style_and_color(self):
        def change(d):
            style = d.styles.add_style('CustomSubheading', 1); style.base_style = d.styles['Heading 2']
            next(p for p in d.paragraphs if p.text == '2.1 技术机制').style = style
        self.check(self.mutate(change))

    def test_theme_color_override_fails(self):
        def change(d):
            color = d.styles['Heading 2'].element.find('w:rPr/w:color', NS)
            color.set(qn('w:themeColor'), 'accent1')
        self.check(self.mutate(change), 'palette mismatch')

    def test_character_style_override_fails(self):
        def change(d):
            style = d.styles.add_style('BlackCharacter', 2); style.font.color.rgb = RGBColor(0, 0, 0)
            next(p for p in d.paragraphs if p.text == '2.1 技术机制').runs[0].style = style
        self.check(self.mutate(change), 'palette mismatch')

    def test_literal_square_fails(self):
        self.check(self.mutate(lambda d: setattr(next(p for p in d.paragraphs if p.text == '2.1 技术机制'), 'text', '▪ 2.1 技术机制')), 'unwanted literal')

    def test_inherited_bullet_fails_and_explicit_cancel_passes(self):
        def change(d, cancel=False):
            # Link Heading 2 to a real bullet numbering definition.
            d.styles['Heading 2'].base_style = d.styles['List Bullet']
            if cancel:
                for p in d.paragraphs:
                    if p.style.name == 'Heading 2':
                        num = p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId(); num.val = 0
        self.check(self.mutate(change), 'inherited list bullets')
        self.check(self.mutate(lambda d: change(d, True)))

    def test_keep_next_remains_allowed(self):
        self.check(self.mutate(lambda d: setattr(d.styles['Heading 2'].paragraph_format, 'keep_with_next', True)))

    def test_broken_link(self):
        def change(d):
            link = d.element.body.find('.//w:hyperlink', NS); link.set(qn('w:anchor'), 'missing')
        self.check(self.mutate(change), 'broken/ambiguous')

    def test_toc_chapter_label_matches_destination(self):
        def change(d):
            node = d.element.body.find('.//w:hyperlink/w:r/w:t', NS)
            node.text = '02 投资概要 '
        self.check(self.mutate(change), 'chapter label differs from target')

    def test_missing_cache_not_confused_with_year_in_title(self):
        def change(d):
            fld = d.element.body.find('.//w:fldSimple', NS)
            fld.find('.//w:t', NS).text = ''
        self.check(self.mutate(change), 'missing cached page')

    def test_pageref_and_link_must_agree(self):
        def change(d):
            fld = d.element.body.find('.//w:fldSimple', NS); fld.set(qn('w:instr'), 'PAGEREF wrong_target \\h')
        self.check(self.mutate(change), 'PAGEREF target differs')

    def test_missing_and_duplicate_toc_entries(self):
        def remove(d):
            entry = next(p for p in d.paragraphs if p.style.name == 'ReportTOC'); entry._p.getparent().remove(entry._p)
        self.check(self.mutate(remove), 'absent from contents')
        def duplicate(d):
            entry = next(p for p in d.paragraphs if p.style.name == 'ReportTOC'); entry._p.addnext(copy.deepcopy(entry._p))
        self.check(self.mutate(duplicate), 'duplicate or non-heading')

    def test_contents_must_not_use_heading_style(self):
        def change(d):
            d.styles['ReportTOC'].base_style = d.styles['Heading 1']
        self.check(self.mutate(change), 'contents entry uses a native heading')

    def test_toc_content_control_is_excluded(self):
        def change(d):
            entries = [p._p for p in d.paragraphs if p.style.name == 'ReportTOC']
            wrapper = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr')
            obj = OxmlElement('w:docPartObj'); gallery = OxmlElement('w:docPartGallery'); gallery.set(qn('w:val'), 'Table of Contents')
            obj.append(gallery); pr.append(obj); wrapper.append(pr)
            content = OxmlElement('w:sdtContent'); wrapper.append(content); entries[0].addprevious(wrapper)
            for el in entries: content.append(el)
        path = self.mutate(change)
        self.check(path)
        self.assertFalse(any(p.text.startswith('01 ') for p in parse_docx(path)))

    def test_complex_toc_field_excludes_unstyled_result_paragraphs(self):
        def change(d):
            h = next(p for p in d.paragraphs if p.text == '1. 投资概要')
            begin = h.insert_paragraph_before()
            run = OxmlElement('w:r'); node = OxmlElement('w:fldChar'); node.set(qn('w:fldCharType'), 'begin'); run.append(node); begin._p.append(run)
            for fragment in (' TO', 'C \\o "1-3" '):
                run = OxmlElement('w:r'); node = OxmlElement('w:instrText'); node.text = fragment; run.append(node); begin._p.append(run)
            h.insert_paragraph_before('Appendix｜Sources / 来源索引')
            h.insert_paragraph_before('16. 目录结果中的章节标签')
            end = h.insert_paragraph_before()
            run = OxmlElement('w:r'); node = OxmlElement('w:fldChar'); node.set(qn('w:fldCharType'), 'end'); run.append(node); end._p.append(run)
        path = self.mutate(change)
        self.assertEqual(sum(p.text.startswith('Appendix｜Sources') for p in parse_docx(path)), 1)
        self.assertFalse(any('目录结果' in p.text for p in parse_docx(path)))

    def test_body_page_cross_reference_is_not_contents(self):
        def change(d):
            p = next(p for p in d.paragraphs if p.text == '2.1 技术机制').insert_paragraph_before('详见业务机制所在页：')
            link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'report_heading_1')
            fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGEREF report_heading_1 \\h')
            run = OxmlElement('w:r'); node = OxmlElement('w:t'); node.text = '3'; run.append(node); fld.append(run); link.append(fld); p._p.append(link)
        path = self.mutate(change)
        self.assertTrue(any('详见业务机制所在页' in p.text for p in parse_docx(path)))
        self.check(path)

    def quadrant(self, data, remove_image=False):
        paras = parse_docx(self.baseline)
        if remove_image:
            for p in paras:
                p.image_alts = [a for a in p.image_alts if not a.startswith('CHART_COMPETITION_QUADRANT:')]
        path = self.root / 'manifest.json'; path.write_text(json.dumps(data), encoding='utf-8')
        audit = Audit(); audit_competition_quadrant(paras, chapter_ranges(paras), path, audit)
        return audit

    def valid_quadrant(self):
        return {'charts': [{
            'alt_text': 'CHART_COMPETITION_QUADRANT: test', 'source_ids': ['S1'],
            'positioning_mode': 'qualitative', 'target_name': 'Target',
            'axes': {a: {'label': a, 'definition': 'capability', 'scale': 'low/high', 'split_basis': 'disclosed boundary'} for a in ('x', 'y')},
            'positioning_evidence': [{'name': n, 'basis': 'source-supported capability', 'source_ids': ['S1']} for n in ('Target', 'Rival')]
        }]}

    def test_quantitative_and_qualitative_evidence(self):
        for mode in ('qualitative', 'quantitative'):
            data = self.valid_quadrant(); data['charts'][0]['positioning_mode'] = mode
            self.assertEqual(self.quadrant(data).failures, [])

    def test_missing_axis_source_or_target_fails(self):
        for mutate in (lambda c: c['axes']['x'].pop('split_basis'),
                       lambda c: c['positioning_evidence'][1].update(source_ids=['UNKNOWN']),
                       lambda c: c.update(target_name='Not in sample')):
            data = self.valid_quadrant(); mutate(data['charts'][0])
            self.assertTrue(self.quadrant(data).failures)

    def test_missing_quadrant_requires_evidenced_skip(self):
        self.assertTrue(self.quadrant({'charts': []}, True).failures)
        data = {'charts': [], 'coverage': [{'chart_id': '17', 'status': 'skipped', 'reason': 'No defensible comparable axis after checking product disclosures.'}]}
        a = self.quadrant(data, True); self.assertEqual(a.failures, []); self.assertTrue(a.warnings)
        self.assertTrue(self.quadrant(data).failures)


if __name__ == '__main__':
    unittest.main()
