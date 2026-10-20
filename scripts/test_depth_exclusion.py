#!/usr/bin/env python3
from __future__ import annotations

import tempfile
from pathlib import Path

import depth_audit
from docx import Document


def main() -> int:
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "research.md"
        path.write_text(
            "<!-- DEPTH_EXCLUDE_START: RESEARCH_BACKSTAGE -->\n"
            "## 渠道覆盖矩阵\n" + ("后台日志" * 1000) + "\n"
            "## 研究对象覆盖矩阵\n" + ("对象记录" * 1000) + "\n"
            "<!-- DEPTH_EXCLUDE_END: RESEARCH_BACKSTAGE -->\n"
            "## 2. 行业纵向叙事\n"
            "正文分析 2026 年收入增长 20%。\n",
            encoding="utf-8",
        )
        chapters, _, _ = depth_audit.split_sections(depth_audit.parse_md(path))
        assert "渠道覆盖矩阵" in chapters[0]["text"]
        assert chapters[2]["chars"] < 100, chapters[2]["chars"]
        assert "后台日志" not in chapters[2]["text"]
        final = Path(td) / 'report.docx'
        doc = Document()
        doc.add_heading('16. 关键尽调问题', level=1)
        doc.add_paragraph('正文分析 2026 年收入增长 20%。')
        doc.add_heading('附录 A 可比公司', level=1)
        doc.add_paragraph('可比公司正文。')
        doc.add_heading('Appendix｜Methods / 定义与方法', level=1)
        doc.add_paragraph('详细计算与会计对账' * 1000)
        doc.add_heading('Appendix｜Sources / 来源索引', level=1)
        doc.add_paragraph('原始资料' * 1000)
        doc.save(final)
        chapters, appendices, _ = depth_audit.split_sections(depth_audit.parse_docx(final))
        assert chapters[16]['chars'] < 100, chapters[16]['chars']
        assert appendices['A']['chars'] < 100, appendices['A']['chars']
        assert '详细计算' not in appendices['A']['text']
        # Introductory chapters must not be pushed toward artificial numerics.
        prose = '先定义概念，再解释工作原理与业务关系。' * 50
        sample = {n: {'chars': len(prose), 'text': prose, 'tables': 1, 'qtitles': 0} for n in (2, 3, 5)}
        _, warnings, _, _ = depth_audit.audit(sample, {}, set(), 'research', research_dir=td)
        assert not any(w.startswith(('§2 量化', '§3 量化')) for w in warnings), warnings
        assert any(w.startswith('§5 量化') for w in warnings), warnings
        # Floors and table requirements follow content, not its old number.
        for technical, history in ((2, 3), (3, 2)):
            blocks = {
                technical: [('p', f'{technical}. 核心技术入门'), ('p', f'{technical}.0 最小知识'),
                            ('p', '技术说明' * 600), ('p', '术语表')],
                history: [('p', f'{history}. 行业与公司历史'), ('p', '历史背景' * 725), ('tbl', '年份 事件')],
            }
            items = [item for n in sorted(blocks) for item in blocks[n]]
            chapters, _, subs = depth_audit.split_sections(items)
            fails, _, _, _ = depth_audit.audit(chapters, {}, subs, 'final')
            assert not any(f.startswith((f'§{technical} ', f'§{history} ')) for f in fails), fails
            chapters[history]['chars'] = 2300
            fails, _, _, _ = depth_audit.audit(chapters, {}, subs, 'final')
            assert any(f.startswith(f'§{history} 字数') for f in fails), fails
            chapters[history]['chars'] = 3000
            chapters[history]['tables'] = 0
            fails, _, _, _ = depth_audit.audit(chapters, {}, subs, 'final')
            assert any(f.startswith(f'§{history} 缺表') for f in fails), fails
            # Research-stage placement checks also use the same roles.
            chapters[technical]['text'] += '\n指引演变与可信度'
            chapters[history]['text'] += '\n术语表'
            fails, _, _, _ = depth_audit.audit(chapters, {}, subs, 'research', research_dir=td)
            assert 'guidance 演变不得放技术入门，统一移至 §9' in fails
            assert any('历史章含重复术语表' in f for f in fails)
    print("DEPTH EXCLUSION TEST: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
